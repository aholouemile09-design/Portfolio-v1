# -*- coding: utf-8 -*-
"""Trois CV ciblés (IA/ML, génie électrique, gestion de projet), en Word puis en PDF.

Les faits viennent tous des CV d'Emile et de ses corrections du 27/09/2026.
Le ton suit ses propres versions « humain » : phrases directes, verbes simples,
un but concret par puce, pas de point final aux puces.

Normes canadiennes : pas de photo ni d'âge, ordre chronologique inversé,
statut de résident permanent, une à deux pages.

    python cv/generer_cv.py
Candidatures (avec téléphone) : cv/*.docx et cv/*.pdf. Site (sans téléphone) : public/cv/*.pdf.
"""

import os
import subprocess
from docx import Document
from docx.enum.text import WD_TAB_ALIGNMENT
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Inches, Pt, RGBColor

ICI = os.path.dirname(os.path.abspath(__file__))
PDF = os.path.join(os.path.dirname(ICI), "public", "cv")
GRIS = RGBColor(0x55, 0x55, 0x55)

NOM = "Kossi Edem Emile D. Aholou"
# Le numéro vit dans cv/telephone.txt, exclu de git : le dépôt est public.
# Il est absent des CV publiés sur le site (choix d'Emile, 27/09/2026).
_TEL_FICHIER = os.path.join(ICI, "telephone.txt")
TEL = open(_TEL_FICHIER, encoding="utf-8").read().strip() if os.path.exists(_TEL_FICHIER) else ""
CONTACT = ("Edmonton, Alberta, Canada  •  " + TEL + "  •  aholou.emile09@gmail.com  •  "
           "linkedin.com/in/kossi-edem-emile-d-aholou  •  oltavia.com")
STATUT = ("Permanent resident of Canada  •  Bilingual English and French  •  "
          "Valid driver's licence  •  Open to relocation")

# ------------------------------------------------------------------ expériences
XODYIA = ("Founder", "Xodyia", "Edmonton, Alberta", "September 2026 – Present")
RSA = ("Development Agent", "Réseau Santé Alberta", "Edmonton, Alberta", "October 2024 – Present")
DESCO = ("Construction Project Manager", "Desco Architectural Firm", "Lomé, Togo", "January 2023 – April 2024")
CH2000 = ("Engineering Manager, Energy Transport & Distribution", "CH2000",
          "Conglomerate Horizon 2000  •  Lomé, Togo", "January 2021 – December 2022")
ARESS = ("Project Coordinator, Solar Energy", "ARESS Togo, solar subsidiary of CH2000", "Lomé, Togo",
         "January 2019 – December 2020")

DIPLOME = ("Engineering Diploma, Electrical Engineering",
           "Institut de Formation Technique Supérieure (IFTS), Togo", "2018",
           "WES: equivalent to a Canadian bachelor's and master's degree")
ATHABASCA = ("Coursework, B.Sc. Computing and Information Systems", "Athabasca University",
             "November 2025 – June 2026", "6 courses: Python, data structures, algorithms, linear algebra")
PMP = ("36-Hour PMP Exam Preparation Course", "Coursework aligned with PMI standards", "Completed",
       "Working toward the PMP designation")
BENEVOLAT = ("Volunteer", "Edmonton's Food Bank", "Edmonton, Alberta", "July 2024")

CV = {
    "AI-ML": {
        "titre": "AI Developer  |  Electrical Engineer building practical AI tools",
        "resume": (
            "Electrical engineer with a degree assessed by WES as equal to a bachelor's and master's, and 5+ "
            "years of project management in energy and construction, now building AI tools for small "
            "businesses. Founded Xodyia, where I built an assistant that answers customers only from a "
            "company's own documents and a prospecting CRM that follows Canada's anti-spam law. Also built "
            "CodeGraft, a learning platform with an AI tutor on the Claude API. Comfortable taking a tool from "
            "a real business problem to a tested, working version. Currently following a structured machine "
            "learning plan: Python and data first, then modelling and deployment."
        ),
        "competences": [
            ("Programming", "Python, JavaScript, TypeScript, SQL"),
            ("Web and back end", "Node.js, React, Next.js, Supabase, SQLite, REST APIs"),
            ("AI", "Claude API integration, prompt guardrails, rate limiting, document search (BM25), "
                   "testing answers against a question set"),
            ("Tools", "Git, GitHub, Vercel, Netlify"),
            ("Currently learning", "NumPy, pandas, scikit-learn"),
        ],
        "projets": [
            ("Document-based answering assistant  |  Xodyia", "JavaScript, BM25", "2026", [
                "Built a search engine in plain JavaScript that answers customer questions only from a "
                "business's own documents, in French and English",
                "Set topics the assistant must never answer, such as health questions, so the conversation "
                "goes to a person instead",
                "Wrote 75 test questions, including trick questions, and fixed a flaw that let unrelated "
                "questions get a wrong answer",
            ]),
            ("CodeGraft  |  Online learning platform", "Next.js, React, Supabase, Claude API  •  codegraft.vercel.app",
             "2026", [
                "Built a machine learning course path with exercises and quizzes, plus a 180-question "
                "bilingual PMP practice exam",
                "Added an AI tutor on the Claude API that guides learners through their projects instead of "
                "doing the work for them, with a limit on requests",
                "Set up user accounts, progress tracking and password reset with Supabase",
            ]),
            ("Prospecting CRM  |  Xodyia", "Node.js, SQLite, Python", "2026", [
                "Built a CRM that blocks any contact without a valid legal basis under Canada's anti-spam law "
                "(CASL), and tracks consent expiry and opt-outs",
                "Added Excel list import, duplicate detection and automatic backups",
            ]),
        ],
        "experiences": [
            (XODYIA, [
                "Started an AI and automation company for small businesses and built three bilingual product "
                "demos: appointment booking, document-based answers, follow-up on customer enquiries",
                "Wrote the sales strategy and prospecting plan, and scoped the first client pilots",
            ]),
            (RSA, [
                "Coordinate program schedules, stakeholder communication and progress reporting across "
                "multiple communities in central Alberta",
                "Built simple tracking tools (spreadsheets, logs, status updates) to follow activities and results",
            ]),
            (DESCO, ["Managed construction projects from planning to closeout, including budgets, schedules, "
                     "risks and monthly client reports, and reduced budget overruns by 15%"]),
            (CH2000, ["Led medium and low-voltage power distribution projects with teams of 20 to 30 people, "
                      "and cut completion times by at least 15% by streamlining scheduling and reporting"]),
            (ARESS, ["Coordinated solar energy projects, prepared cost estimates and advised management on "
                     "technical questions"]),
        ],
        "formation": [DIPLOME, ATHABASCA, PMP],
        "benevolat": False,
    },

    "Electrical-Engineering": {
        "titre": "Electrical Engineer  |  Power Distribution (MV/LV) and Solar Projects",
        "resume": (
            "Electrical engineer with a degree assessed by WES as equal to a bachelor's and master's, and 5+ "
            "years in power distribution, solar energy and construction projects. Led medium and low-voltage "
            "transmission and distribution projects from design review to site delivery, with teams of 20 to 30 "
            "people. Cut completion times by at least 15% and raised team productivity by 20%. Comfortable with "
            "AutoCAD Electrical, CANECO and MATLAB, and with keeping contractors, consultants and clients "
            "aligned on quality and safety."
        ),
        "competences": [
            ("Power systems", "MV/LV distribution, solar PV projects, drawing and specification review, "
                              "cost estimates"),
            ("Software", "AutoCAD Electrical, CANECO, MATLAB, LabVIEW, Primavera P6, MS Project, MS Office"),
            ("Project delivery", "Scheduling, budgets, risk management, contractor coordination, quality and "
                                 "compliance"),
            ("Safety", "WHMIS 2015, Construction Safety Training System (CSTS)"),
        ],
        "projets": [],
        "titre_experiences": "Engineering Experience",
        "experiences": [
            (DESCO, [
                "Prepared cost estimates and followed budgets on multi-phase projects, reducing budget "
                "overruns by 15%",
                "Managed contracts, checked compliance and coordinated contractors and consultants to meet "
                "quality standards",
            ]),
            (CH2000, [
                "Led medium and low-voltage power transmission and distribution projects across several sites",
                "Reviewed engineering designs and approved them or asked for changes to improve compliance "
                "and reduce risk",
                "Hired, trained and directed teams of 20 to 30 people and raised productivity by 20%",
                "Cut project completion times by at least 15% by streamlining scheduling and status reporting",
                "Won a follow-up contract with the same client after delivering the first project",
            ]),
            (ARESS, [
                "Prepared cost estimates for upcoming solar projects to support budgeting and resource planning",
                "Advised senior management on technical questions and engineering decisions",
                "Checked project data for accuracy and set up risk management steps for technical work to "
                "improve safety",
            ]),
        ],
        "autres": [
            (XODYIA, ["Started an AI and automation company for small businesses"]),
            (RSA, ["Coordinate program schedules, stakeholder communication and progress reporting across "
                   "central Alberta"]),
        ],
        "formation": [DIPLOME[:3] + (DIPLOME[3] + "; power and low-current systems",), PMP],
        "benevolat": False,
    },

    "Project-Management": {
        "titre": "Project Manager  |  Energy, Construction and Community Programs",
        "resume": (
            "Project manager with an electrical engineering degree assessed by WES as equal to a bachelor's and "
            "master's, and 5+ years of experience delivering energy, construction and community programs in "
            "Togo and Canada. Track record of coordinating stakeholders, preparing project documentation and "
            "status reports, managing risk and keeping budgets and schedules on track. Cut project completion "
            "times by at least 15% and reduced budget overruns by 15%. Working toward the PMP designation."
        ),
        "competences": [
            ("Project coordination", "Project initiation and documentation, scheduling, budgets and cost "
                                     "estimates, risk management, procurement support"),
            ("Reporting and relationships", "Monthly status reports, client and partner relations, consultant "
                                            "and contractor coordination, bilingual communication"),
            ("Tools", "MS Project, Primavera P6, MS Office (Word, Excel, Outlook, PowerPoint)"),
        ],
        "projets": [],
        "experiences": [
            (XODYIA, [
                "Started a new company from incorporation to launch: website requirements and follow-up with the "
                "developer, prospecting plan, product demos and first client pilots",
            ]),
            (RSA, [
                "Coordinate program schedules, stakeholder communication and progress reporting across multiple "
                "communities in central Alberta",
                "Prepare correspondence and documentation for partner review, with limited direction",
            ]),
            (DESCO, [
                "Coordinated stakeholder communication across clients, consultants and internal teams to support "
                "timely decision-making throughout the project lifecycle",
                "Prepared project initiation, planning and closeout documentation, and reviewed vendor proposals "
                "to support procurement decisions",
                "Identified and managed project risks, and prepared monthly status reports for clients",
                "Coordinated contractors and consultants and followed budgets closely, reducing budget overruns "
                "by 15%",
            ]),
            (CH2000, [
                "Directed and coordinated teams of 20 to 30 people across active projects, keeping sites aligned",
                "Reduced project completion times by at least 15% by streamlining scheduling and status reporting",
                "Reviewed technical drawings and specifications to track deliverables against approved plans",
            ]),
            (ARESS, [
                "Coordinated solar energy projects, including cost estimates, resource planning and risk management",
            ]),
        ],
        "formation": [PMP, DIPLOME],
        "benevolat": False,
    },
}


# ------------------------------------------------------------------ mise en forme
def filet(paragraphe):
    """Trait fin sous un titre de section."""
    p = paragraphe._p.get_or_add_pPr()
    bdr = OxmlElement("w:pBdr")
    bas = OxmlElement("w:bottom")
    for k, v in (("val", "single"), ("sz", "6"), ("space", "1"), ("color", "999999")):
        bas.set(qn(f"w:{k}"), v)
    bdr.append(bas)
    p.append(bdr)


def espace(p, avant=0, apres=0):
    p.paragraph_format.space_before = Pt(avant)
    p.paragraph_format.space_after = Pt(apres)


def texte(d, t, taille=10, gras=False, couleur=None, avant=0, apres=2):
    p = d.add_paragraph()
    r = p.add_run(t)
    r.font.size, r.bold = Pt(taille), gras
    if couleur:
        r.font.color.rgb = couleur
    espace(p, avant, apres)
    return p


def section(d, titre):
    p = texte(d, titre.upper(), 10.5, gras=True, avant=6, apres=3)
    p.runs[0].font.spacing = Pt(1)
    filet(p)


def entete(d, gauche, sous, droite):
    """Ligne titre en gras, date alignée à droite, puis une ligne grise."""
    p = d.add_paragraph()
    p.paragraph_format.tab_stops.add_tab_stop(Inches(7.3), WD_TAB_ALIGNMENT.RIGHT)
    r = p.add_run(gauche)
    r.bold, r.font.size = True, Pt(10.5)
    r = p.add_run("\t" + droite)
    r.font.size, r.font.color.rgb = Pt(9.5), GRIS
    espace(p, 5, 0)
    if sous:
        texte(d, sous, 9.5, couleur=GRIS, apres=1)


def puce(d, t):
    p = d.add_paragraph(style="List Bullet")
    p.paragraph_format.left_indent = Inches(0.22)
    r = p.add_run(t)
    r.font.size = Pt(10)
    espace(p, 0, 1)


def postes(d, liste):
    for (role, org, lieu, dates), puces in liste:
        entete(d, f"{role}  |  {org}", lieu, dates)
        for t in puces:
            puce(d, t)


def generer(cle, cv, web=False):
    d = Document()
    for s in d.sections:
        s.top_margin = s.bottom_margin = Inches(0.45)
        s.left_margin = s.right_margin = Inches(0.6)
    st = d.styles["Normal"]
    st.font.name = "Calibri"
    st.element.rPr.rFonts.set(qn("w:eastAsia"), "Calibri")
    st.paragraph_format.line_spacing = 1.08

    texte(d, NOM, 20, gras=True, apres=1)
    texte(d, cv["titre"], 11.5, couleur=GRIS, apres=3)
    texte(d, CONTACT.replace(TEL + "  •  ", "") if web else CONTACT, 9, apres=0)
    texte(d, STATUT, 9, couleur=GRIS, apres=2)

    section(d, "Professional summary")
    texte(d, cv["resume"], 10, apres=0)

    section(d, "Core competencies")
    for titre, valeur in cv["competences"]:
        p = d.add_paragraph()
        r = p.add_run(titre + ": ")
        r.bold, r.font.size = True, Pt(10)
        p.add_run(valeur).font.size = Pt(10)
        espace(p, 0, 1)

    if cv["projets"]:
        section(d, "Selected projects")
        for titre, sous, date, puces in cv["projets"]:
            entete(d, titre, sous, date)
            for t in puces:
                puce(d, t)

    section(d, cv.get("titre_experiences", "Professional experience"))
    postes(d, cv["experiences"])
    if cv.get("autres"):
        section(d, "Other experience")
        postes(d, cv["autres"])

    section(d, "Education & professional development")
    for titre, ecole, date, note in cv["formation"]:
        entete(d, titre, f"{ecole}  •  {note}" if note else ecole, date)

    if cv["benevolat"]:
        section(d, "Volunteering")
        role, org, lieu, date = BENEVOLAT
        entete(d, f"{role}, {org}, {lieu}", "", date)

    chemin = os.path.join(ICI, "_web" if web else "", f"CV_Emile_Aholou_{cle}.docx")
    os.makedirs(os.path.dirname(chemin), exist_ok=True)
    d.save(chemin)
    return chemin


def en_pdf(fichiers, dossier):
    """Conversion par Word lui-même : le PDF est identique au .docx."""
    os.makedirs(dossier, exist_ok=True)
    lignes = ["$w = New-Object -ComObject Word.Application", "$w.Visible = $false"]
    for f in fichiers:
        sortie = os.path.join(dossier, os.path.basename(f).replace(".docx", ".pdf"))
        lignes += [f"$d = $w.Documents.Open('{f}', $false, $true)",
                   f"$d.SaveAs2('{sortie}', 17)",
                   "Write-Output ('pages ' + $d.ComputeStatistics(2) + ' : ' + $d.Name)",
                   "$d.Close($false)"]
    lignes.append("$w.Quit()")
    r = subprocess.run(["powershell", "-NoProfile", "-Command", "; ".join(lignes)],
                       capture_output=True, text=True)
    print(r.stdout.strip() or r.stderr.strip())


if __name__ == "__main__":
    # Candidatures : avec téléphone, .docx et .pdf dans cv/.
    en_pdf([generer(cle, cv) for cle, cv in CV.items()], ICI)
    # Site : sans téléphone, PDF dans public/cv/.
    en_pdf([generer(cle, cv, web=True) for cle, cv in CV.items()], PDF)
    print("PDF :", PDF)
